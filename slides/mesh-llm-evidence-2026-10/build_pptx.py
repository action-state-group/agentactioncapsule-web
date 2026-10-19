#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""Build the two-slide Mesh-LLM evidence deck as native, editable PowerPoint shapes.

Re-run:  python3 build_pptx.py [out.pptx]   (needs python-pptx, pymupdf, lxml; Liberation Sans
         from LIBERATION_DIR, default /usr/share/fonts/truetype/liberation)
All slide text lives in TEXT below. Layout is computed from measured text (Liberation Sans ==
Arial metrics); the script raises if any text would not fit its box at the requested size.
"""
import os
import sys

import pymupdf
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ----------------------------------------------------------------------------------------------
# ALL TEXT (edit here only). "\n" inside a string = line break within the same paragraph.
# Box / band items are (bold lead, rest) pairs so the term before the colon can be bold.
# ----------------------------------------------------------------------------------------------
TEXT = {
    "s1": {
        "title": "Mesh-LLM Evidence Layer:\nWhat Trust between Strangers Looks Like",
        "subtitle": "What a node says it served, and its history of what it has done "
                    "(targeting mesh-llm v0.79)",
        "diagram": {
            "node_a": "Node A", "node_b": "Node B",
            "request": "request", "answer": "answer",
            "log_a": "A's log", "log_b": "B's log",
            "capsule": "capsule", "cites": "cites",
            "checkpoint": "checkpoint",
            "witness": "witness\n(optional)",
            "report": "report",
            "verifier": "outside\nverifier",
        },
        "boxes": [
            ("Capsule:", " a signed record of one action. Digests, never the prompt or answer."),
            ("Log + checkpoint:", " append-only and signed, with no central server."),
            ("Bilateral:", " both sides record the same exchange, so any gap or mismatch shows."),
            ("Witness:", " an outside timestamp. It shows when, not what."),
            ("Third-party report:", " anyone checks it offline, with no account."),
        ],
        "italic": "“The mesh needs no central authority to trust: each node keeps its own "
                  "receipt and checks the other's.”",
        "notes": [
            "The mesh needs no central authority to trust: each node keeps its own receipt and "
            "checks the other's.",
            "One: a capsule is a signed record of one action; it holds digests, never the prompt "
            "or the answer.",
            "Two: each node keeps a log with checkpoints that is append-only and signed, with no "
            "central server.",
            "Three: bilateral means both sides record the same exchange, so any gap or mismatch "
            "shows.",
            "Four: if the operator names a witness, it adds an outside timestamp; it shows when, not what.",
            "Five: the third-party report is something anyone can check offline, with no account; for a mesh node's records, that offline check is not yet joined end to end.",
        ],
    },
    "s2": {
        "title": "What each check shows, and what it doesn't",
        "header": ["Question", "Check", "Shows", "Does not show"],
        "rows": [
            ["Who sent this, and was it changed?", "Signed",
             "who signed; unchanged since", "that the content is true"],
            ["Was it there at the time?", "Checkpointed",
             "in the log by time T, in order", "that everything was logged"],
            ["Did we both see the same exchange?", "Both sides match",
             "same request and answer, cited both ways", "which side is right"],
            ["Was the claimed model served?", "Second node + referee",
             "a signed ruling on one client-marked pair",
             "every request; the referee is off by default today"],
            ["What actually ran?", "Hardware attestation (next)",
             "a measured runtime", "more than the quote covers"],
        ],
        "band_title": "What keeps everyone honest",
        "band_items": [
            "each side's own signed log",
            "the other side's copy",
            "an outside witness",
            "a referee neither side picked",
        ],
        "arrow": "builds a history per peer (counts, never a score)",
        "deeper_lead": "Go deeper: ",
        "deeper": "agentactioncapsule.org · datatracker.ietf.org/doc/search?name=draft-mih "
                  "(AAC, checkpointed log, bilateral, evidence layer; individual Internet-Drafts) "
                  "· SCITT, RFC 9943 · github.com/Mesh-LLM/capsules (evidence plugin)",
        "italic": "“No reputation score anywhere. You get a map of what was checked, and the "
                  "gaps are named.”",
        "notes": [
            "No reputation score anywhere. You get a map of what was checked, and the gaps are "
            "named.",
            "Who sent this, and was it changed? A signature shows who signed and that it is "
            "unchanged since; it does not show that the content is true.",
            "Was it there at the time? A checkpoint shows it was in the log by time T, in order; "
            "it does not show that everything was logged.",
            "Did we both see the same exchange? When both sides match, that shows the same request "
            "and answer, cited both ways; it does not show which side is right.",
            "Was the claimed model served? When a second node answers a request the client marked and the two answers differ, "
            "a referee neither side picked signs a ruling on that one pair; it does not cover "
            "every request, and the referee is off by default today.",
            "What actually ran? Hardware attestation, which comes next, shows a measured runtime; "
            "it does not show more than the quote covers.",
        ],
    },
}

# ----------------------------------------------------------------------------------------------
# Style (STYLE.md). Units are px on a 1920x1080 stage; 1px = 9525 EMU; pt = px * 0.75.
# ----------------------------------------------------------------------------------------------
C = {k: RGBColor.from_string(v) for k, v in {
    "bg": "FFFFFF", "ink": "111827", "muted": "374151", "line": "9CA3AF", "panel": "F3F4F6",
    "accent": "1D4ED8", "accent2": "0F766E", "gap": "B45309", "gapbg": "FFF7ED",
    "white": "FFFFFF"}.items()}
FONT = "Arial"
TITLE_PX, SUB_PX, BODY_PX, LABEL_PX, DEEPER_PX = 64, 40, 38, 30, 38
BORDER_PX, RADIUS_PX = 3, 16
ML, MR, MT, MB = 80, 80, 64, 56
W, H = 1920, 1080
LINE_H = 1.15          # rendered line height / font size for Arial single spacing (measured margin)
FIT_SLACK = 0.97       # use only 97% of the box width when wrapping (renderer safety)

_HERE = os.path.dirname(os.path.abspath(__file__))
_FONTDIR = os.environ.get("LIBERATION_DIR", "/usr/share/fonts/truetype/liberation")
_FONTS = {(b, i): pymupdf.Font(fontfile=os.path.join(_FONTDIR, f"LiberationSans-{n}.ttf"))
          for (b, i), n in {(False, False): "Regular", (True, False): "Bold",
                            (False, True): "Italic", (True, True): "BoldItalic"}.items()}


def px(v):
    return Emu(int(round(v * 9525)))


# ---------------------------------------------------------------- measuring
def wrap_lines(runs, width, size, slack=FIT_SLACK):
    """runs: [(text, bold, italic)]. Greedy word wrap like PowerPoint; returns number of lines."""
    width *= slack
    lines = 0
    # split runs on hard line breaks
    segs, cur = [], []
    for text, b, i in runs:
        parts = text.split("\n")
        for k, p in enumerate(parts):
            if k:
                segs.append(cur)
                cur = []
            cur.append((p, b, i))
    segs.append(cur)
    for seg in segs:
        # tokens: words with their trailing space, carrying font
        toks = []
        for text, b, i in seg:
            f = _FONTS[(b, i)]
            for w in text.split(" "):
                toks.append((f.text_length(w, size), f.text_length(" ", size)))
        x, n = 0.0, 1
        for wl, sl in toks:
            if x > 0 and x + wl > width:
                n += 1
                x = 0.0
            if wl > width:
                raise ValueError(f"word wider than box ({wl:.0f} > {width:.0f}) in {seg}")
            x += wl + sl
        lines += n
    return lines


def hard_wrap(runs, width, size, slack=FIT_SLACK):
    """Insert explicit line breaks where our greedy wrap breaks, so no renderer ever breaks a
    URL at '-' or '?'. Only spaces become breaks; the visible text is unchanged."""
    width *= slack
    out, x = [], 0.0
    for text, b, i in runs:
        f = _FONTS[(b, i)]
        words = text.split(" ")
        buf = ""
        for k, w in enumerate(words):
            wl, sl = f.text_length(w, size), f.text_length(" ", size)
            if k:
                if x > 0 and x + wl > width:
                    buf += "\n"
                    x = 0.0
                else:
                    buf += " "
            buf += w
            x += wl + sl
        out.append((buf, b, i))
    return out


def text_h(runs, width, size, slack=FIT_SLACK):
    return wrap_lines(runs, width, size, slack) * size * LINE_H


# ---------------------------------------------------------------- shape helpers
def _line(shape_line, color, width_px=BORDER_PX, dash=None):
    if color is None:
        shape_line.fill.background()
        return
    shape_line.color.rgb = C[color]
    shape_line.width = px(width_px)
    if dash:
        shape_line.dash_style = dash


def _drop_style(shp):
    """Remove the theme style reference (it pulls in theme shadows/bevels). Fill, line and text
    colour are all set explicitly, so the shape renders the same in PowerPoint, Keynote, Slides."""
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)


def box(slide, kind, x, y, w, h, fill=None, line="line", dash=None, radius=RADIUS_PX,
        line_w=BORDER_PX):
    shp = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius / min(w, h)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = C[fill]
    else:
        shp.fill.background()
    _line(shp.line, line, line_w, dash)
    _drop_style(shp)
    tf = shp.text_frame
    tf.text = ""
    return shp


def fill_text(shp, paras, size, color="ink", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
              inset=(0, 0, 0, 0), check=True):
    """paras: list of paragraphs; each paragraph = list of (text, bold, italic) runs."""
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    l, t, r, b = inset
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = px(l), px(t), px(r), px(b)
    tf.vertical_anchor = anchor
    if check:
        avail_w = shp.width / 9525 - l - r
        need = sum(text_h(p, avail_w, size) for p in paras)
        avail_h = shp.height / 9525 - t - b
        if need > avail_h + 0.5:
            raise ValueError(f"text overflows ({need:.0f}px > {avail_h:.0f}px): {paras}")
    for k, runs in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.0
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for text, bold, italic in runs:
            pieces = text.split("\n")
            for j, piece in enumerate(pieces):
                if j:
                    p.add_line_break()
                if not piece:
                    continue
                run = p.add_run()
                run.text = piece
                f = run.font
                f.name = FONT
                f.size = Pt(size * 0.75)
                f.bold = bold
                f.italic = italic
                f.color.rgb = C[color]
                # pin the east-asian / complex-script faces too, so nothing falls back
                rPr = run._r.get_or_add_rPr()
                for tag in ("a:ea", "a:cs"):
                    el = rPr.find(qn(tag))
                    if el is None:
                        el = etree.SubElement(rPr, qn(tag))
                    el.set("typeface", FONT)
    return shp


def textbox(slide, x, y, w, h, paras, size, **kw):
    shp = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    return fill_text(shp, paras, size, **kw)


def connector(slide, x1, y1, x2, y2, color="accent", width_px=BORDER_PX, dash=None,
              head=False, tail=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    _line(c.line, color, width_px, dash)
    _drop_style(c)
    ln = c.line._get_or_add_ln()
    for want, tag in ((head, "a:headEnd"), (tail, "a:tailEnd")):
        if want:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", "triangle")
            e.set("w", "med")
            e.set("len", "med")
    return c


def R(text, bold=False, italic=False):
    return [(text, bold, italic)]


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C["bg"]
    return s


def notes(slide, lines):
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def italic_line(slide, text):
    # full-width line: measured exactly (no slack); the 1-line fit is confirmed in the render
    h = text_h([(text, False, True)], W - ML - MR + 30, BODY_PX, slack=1.0)
    y = H - MB - h
    textbox(slide, ML - 15, y, W - ML - MR + 30, h, [R(text, italic=True)], BODY_PX,
            align=PP_ALIGN.CENTER, check=False)
    return y


# ---------------------------------------------------------------- slide 1
def slide1(prs):
    T = TEXT["s1"]
    D = T["diagram"]
    s = new_slide(prs)
    cw = W - ML - MR
    th = text_h(R(T["title"], True), cw, TITLE_PX)
    textbox(s, ML, MT - 16, cw, th, [R(T["title"], True)], TITLE_PX)
    sy = MT - 16 + th + 8
    sh = text_h(R(T["subtitle"]), cw, SUB_PX)
    textbox(s, ML, sy, cw, sh, [R(T["subtitle"])], SUB_PX, color="muted")
    top = sy + sh + 26
    bottom = italic_line(s, T["italic"]) - 24

    # ---- right column: five numbered boxes, joined by arrows
    bx, bw = 900, W - MR - 900
    badge = 60
    pad = 10
    text_x_off = 18 + badge + 18
    tw = bw - text_x_off - 20
    hs = [max(text_h([(a, True, False), (b, False, False)], tw, BODY_PX) + 2 * pad, badge + 2 * pad)
          for a, b in T["boxes"]]
    gap = (bottom - top - sum(hs)) / (len(hs) - 1)
    if gap < 26:
        raise ValueError(f"boxes do not fit vertically (gap {gap:.0f})")
    gap = min(gap, 48)
    y = top + (bottom - top - sum(hs) - gap * (len(hs) - 1)) / 2
    for k, ((lead, rest), h) in enumerate(zip(T["boxes"], hs)):
        b = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx, y, bw, h, fill="panel")
        fill_text(b, [[(lead, True, False), (rest, False, False)]], BODY_PX,
                  inset=(text_x_off, pad, 20, pad), anchor=MSO_ANCHOR.MIDDLE)
        o = box(s, MSO_SHAPE.OVAL, bx + 18, y + (h - badge) / 2, badge, badge, fill="accent",
                line=None)
        fill_text(o, [R(str(k + 1), True)], BODY_PX, color="white", align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)
        if k:
            cx = bx + 18 + badge / 2
            connector(s, cx, y - gap + 2, cx, y - 2)
        y += h + gap

    # ---- left: compact diagram in [ML, 840] x [top, bottom]
    x0, x1 = ML, 840
    nw, nh = 230, 96
    ax, bxn = x0, x1 - nw          # node A / node B left edges
    ny = top + 8
    for nx, label in ((ax, D["node_a"]), (bxn, D["node_b"])):
        n = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, nx, ny, nw, nh, fill="white", line="ink")
        fill_text(n, [R(label, True)], BODY_PX, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # request A->B, answer B->A
    ry, ay = ny + 30, ny + nh - 30
    connector(s, ax + nw + 6, ry, bxn - 6, ry)
    connector(s, bxn - 6, ay, ax + nw + 6, ay)
    mid_l, mid_w = ax + nw, bxn - ax - nw
    lh = LABEL_PX * LINE_H
    textbox(s, mid_l, ry - lh - 4, mid_w, lh, [R(D["request"])], LABEL_PX, color="muted",
            align=PP_ALIGN.CENTER)
    textbox(s, mid_l, ay + 4, mid_w, lh, [R(D["answer"])], LABEL_PX, color="muted",
            align=PP_ALIGN.CENTER)

    # logs below each node, each with a capsule and a checkpoint mark
    ly = ny + nh + 44
    lw, lhgt = nw, 250
    cap_h = 58
    cap_y = ly + 62
    cp_y = cap_y + cap_h + 34         # checkpoint rule
    caps = []
    for lx, label in ((ax, D["log_a"]), (bxn, D["log_b"])):
        connector(s, lx + nw / 2, ny + nh + 4, lx + nw / 2, ly - 4)       # node drops capsule
        lg = box(s, MSO_SHAPE.RECTANGLE, lx, ly, lw, lhgt, fill="panel")
        fill_text(lg, [R(label, True)], BODY_PX, align=PP_ALIGN.CENTER, inset=(0, 8, 0, 0))
        cp = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, lx + 25, cap_y, lw - 50, cap_h, fill="white",
                 line="accent")
        fill_text(cp, [R(D["capsule"])], BODY_PX, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        caps.append(cp)
        # checkpoint: accent rule across the log + diamond mark + label
        connector(s, lx + 12, cp_y, lx + lw - 12, cp_y, color="accent", tail=False)
        dm = 26
        box(s, MSO_SHAPE.DIAMOND, lx + 14, cp_y - dm / 2, dm, dm, fill="accent", line=None)
        textbox(s, lx + 8, cp_y + 10, lw - 16, lh, [R(D["checkpoint"])], LABEL_PX,
                color="accent", align=PP_ALIGN.CENTER)
    # dashed "cites" link between the two capsules, both directions
    cyc = cap_y + cap_h / 2
    connector(s, ax + lw - 25 + 6, cyc, bxn + 25 - 6, cyc, color="ink", dash=MSO_LINE.DASH,
              head=True, tail=True, width_px=3)
    textbox(s, mid_l, cyc - lh - 6, mid_w, lh, [R(D["cites"])], LABEL_PX, color="muted",
            align=PP_ALIGN.CENTER)

    # optional witness: dashed outline, between the logs, tied to both checkpoints
    ww, wh = 210, text_h(R(D["witness"]), 210, BODY_PX) + 16
    wx, wy = (x0 + x1) / 2 - ww / 2, cp_y - wh / 2 + 18
    wbox = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, wx, wy, ww, wh, fill="white", line="accent2",
               dash=MSO_LINE.DASH)
    fill_text(wbox, [R(D["witness"], True)], BODY_PX, color="accent2", align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, inset=(0, 8, 0, 8))
    connector(s, ax + lw + 4, cp_y, wx - 4, cp_y, color="accent2", dash=MSO_LINE.DASH,
              head=True, tail=False)
    connector(s, wx + ww + 4, cp_y, bxn - 4, cp_y, color="accent2", dash=MSO_LINE.DASH,
              head=False, tail=True)

    # report, read by an outside verifier
    rw, rh = 220, 100
    ry2 = bottom - rh - 6
    rx = x0 + 190
    if ry2 < ly + lhgt + 40:
        raise ValueError("diagram bottom row does not fit")
    rp = box(s, MSO_SHAPE.FLOWCHART_DOCUMENT, rx, ry2, rw, rh, fill="white", line="ink")
    fill_text(rp, [R(D["report"], True)], BODY_PX, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, inset=(0, 0, 0, 14))
    vw = 230
    vh = text_h(R(D["verifier"]), vw, BODY_PX) + 16
    vx, vy = x1 - vw, ry2 + rh / 2 - vh / 2
    vb = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, vx, vy, vw, vh, fill="white", line="accent2")
    fill_text(vb, [R(D["verifier"], True)], BODY_PX, color="accent2", align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, inset=(0, 8, 0, 8))
    # logs feed the report; report goes to the verifier
    connector(s, ax + lw / 2, ly + lhgt + 4, rx + rw / 2, ry2 - 4)
    connector(s, bxn + lw / 2, ly + lhgt + 4, rx + rw - 30, ry2 - 4)
    connector(s, rx + rw + 6, ry2 + rh / 2 - 6, vx - 6, ry2 + rh / 2 - 6)

    notes(s, T["notes"])


# ---------------------------------------------------------------- slide 2
def slide2(prs):
    T = TEXT["s2"]
    s = new_slide(prs)
    cw = W - ML - MR
    th = text_h(R(T["title"], True), cw, TITLE_PX)
    ty0 = MT - 24
    textbox(s, ML, ty0, cw, th, [R(T["title"], True)], TITLE_PX)
    top = ty0 + th + 18
    iy = italic_line(s, T["italic"])

    # ---- trust map: a grid of native rectangles, full width, every cell <= 2 lines
    pad_x, pad_y, rule = 12, 7, 8
    inner = lambda j: 2 * pad_x + (rule if j == 3 else 0)
    cols = []
    for j in range(4):
        cells = [r[j] for r in T["rows"]] + [T["header"][j]]
        w = 150
        while True:
            try:
                if max(wrap_lines(R(t, bold=(j == 1)), w - inner(j), BODY_PX) for t in cells) <= 2:
                    break
            except ValueError:
                pass
            w += 5
        cols.append(w)
    spare = cw - sum(cols)
    if spare < 0:
        raise ValueError(f"table wider than slide by {-spare}px")
    cols = [c + spare * c / sum(cols) for c in cols]
    hdr_h = BODY_PX * LINE_H + 2 * pad_y + 4
    row_h = [max(text_h(R(t, bold=(j == 1)), cols[j] - inner(j), BODY_PX) for j, t in enumerate(r))
             + 2 * pad_y for r in T["rows"]]
    total = hdr_h + sum(row_h)

    styles = [("panel", "ink"), ("panel", "ink"), ("accent", "white"), ("gap", "white")]
    x = ML
    for j, (label, (fill, ink)) in enumerate(zip(T["header"], styles)):
        c = box(s, MSO_SHAPE.RECTANGLE, x, top, cols[j], hdr_h, fill=fill, line="line")
        fill_text(c, [R(label, True)], BODY_PX, color=ink, anchor=MSO_ANCHOR.MIDDLE,
                  inset=(inner(j) - pad_x, pad_y, pad_x, pad_y))
        x += cols[j]
    y = top + hdr_h
    for row, h in zip(T["rows"], row_h):
        x = ML
        for j, t in enumerate(row):
            gapcol = j == 3
            c = box(s, MSO_SHAPE.RECTANGLE, x, y, cols[j], h,
                    fill="gapbg" if gapcol else "white", line="line")
            fill_text(c, [R(t, bold=(j == 1))], BODY_PX, anchor=MSO_ANCHOR.MIDDLE,
                      inset=(inner(j) - pad_x, pad_y, pad_x, pad_y))
            x += cols[j]
        y += h
    gx = ML + sum(cols[:3])          # "Does not show" left rule
    box(s, MSO_SHAPE.RECTANGLE, gx, top, rule - 1, total, fill="gap", line=None)

    # ---- bottom zone: [arrow line + Go deeper] on the left, side band on the right
    zt, zb = top + total + 18, iy - 16
    zh = zb - zt
    bp = 14
    bparas = [R(T["band_title"], True)] + [R("\u2022 " + it) for it in T["band_items"]]
    band_w = max(_FONTS[(True, False)].text_length(T["band_title"], BODY_PX),
                 *(_FONTS[(False, False)].text_length("\u2022 " + it, BODY_PX)
                   for it in T["band_items"])) / FIT_SLACK + 2 * bp + 6
    band_x = W - MR - band_w
    b = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, band_x, zt, band_w, zh, fill="white",
            line="accent2")
    fill_text(b, bparas, BODY_PX, inset=(bp, 6, bp, 6), anchor=MSO_ANCHOR.MIDDLE)
    for r in b.text_frame.paragraphs[0].runs:
        r.font.color.rgb = C["accent2"]

    lw = band_x - 24 - ML
    ah = BODY_PX * LINE_H
    aw = 100
    box(s, MSO_SHAPE.RIGHT_ARROW, ML, zt + 6, aw, ah - 12, fill="accent", line=None)
    textbox(s, ML + aw + 16, zt, lw - aw - 16, ah, [R(T["arrow"], True)], BODY_PX,
            color="accent", anchor=MSO_ANCHOR.MIDDLE)
    druns = [(T["deeper_lead"], True, False), (T["deeper"], False, False)]
    avail = zb - (zt + ah + 12)
    dpx = DEEPER_PX
    while text_h(druns, lw, dpx) > avail and dpx > 32:
        dpx -= 1
    dh = text_h(druns, lw, dpx)
    textbox(s, ML, zb - dh, lw, dh, [hard_wrap(druns, lw, dpx)], dpx, color="muted")
    if zb - dh < zt + ah + 4:
        raise ValueError("go-deeper block does not fit")

    notes(s, T["notes"])
    return dpx


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    slide1(prs)
    dpx = slide2(prs)
    prs.save(out)
    print(f"wrote {out}; go-deeper size {dpx}px = {dpx * 0.75}pt")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else
         os.path.join(_HERE, "mesh-llm-evidence.pptx"))
